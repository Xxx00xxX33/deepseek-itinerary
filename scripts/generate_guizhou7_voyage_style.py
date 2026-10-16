#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""多彩贵州7天（重庆进出）— Voyage 样式行程，来源：华宇环球旅游4月报价（重庆进出西部航空）.doc。
不含任何价格、餐标数字、供应商信息；景点与服务仅限报价所列。
Day 条只列报价膳食栏的午/晚餐名；文案已按第三方独立审计意见修正（见 references/independent-audit.md）。"""
import datetime
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import generate_chongqing_sanxia8_voyage_style as g  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402

TEMPLATE = r'C:\Users\DELL\Downloads\12D11N_东北草原深度游_2026年6月.docx'
OUTPUT = r'C:\Users\DELL\Downloads\多彩贵州7天（重庆进出）_行程.docx'

TITLE = '7天 多彩贵州·山水苗寨之旅'
SUBTITLE = '新加坡·重庆·安顺·黔西·贵阳·凯里·荔波'
FEATURES = '行程亮点：龙宫·黄果树瀑布·陡坡塘·百里杜鹃·西江千户苗寨·荔波大小七孔·甲秀楼'

HOTEL_AS = '凌悦酒店或西悦臻选酒店或同级（当地准五）'
HOTEL_GY = '美仑酒店或浩枫温德姆酒店或同级（当地准五）'
HOTEL_KL = '千云山房大酒店或龙都璟怡大酒店或同级（当地准五）'
HOTEL_LB = '天泰酒店或嘉和酒店或同级（当地准五）'

SIGHT_INTROS = {
    '龙宫景区（电梯+游船）':
        '藏在安顺群山里的地下水世界，是喀斯特地貌孕育的大型水溶洞。搭乘景区电梯，再坐小船驶入幽暗的地下暗河，'
        '钟乳石在灯光下五彩斑斓，四周静谧幽深，仿佛真的闯进了传说中的“龙王宫殿”。',
    '黄果树瀑布景区（环保车+往返扶梯）':
        '中国最著名的大瀑布之一，高约77.8米、宽约101米。乘环保车入园，搭扶梯往返省下脚力，'
        '还未见瀑布，先闻轰鸣如雷；走近时水雾扑面，晴天时有机会看到彩虹挂在飞瀑之间，气势令人震撼。',
    '陡坡塘瀑布景区':
        '黄果树瀑布群中瀑顶最宽的一座，宽逾百米，水流如一面银色的帘幕铺展开来。'
        '这里正是1986年版电视剧《西游记》片尾师徒四人牵马过瀑的取景地，很多人会在此重温“敢问路在何方”的经典画面。',
    '百里杜鹃景区（电瓶车）':
        '黔西北以杜鹃花闻名的景区，绵延百余里的原生杜鹃林带被誉为“地球彩带”。每年春季花开时节，'
        '红、粉、紫、白的杜鹃次第开放。景区内乘电瓶车游览。花期及盛花程度受当年气候影响，以实际为准。',
    '西江千户苗寨（四段电瓶车+苗服体验+无人机拍摄团体小视频）':
        '全世界规模最大的苗族聚居村寨，千余户木构吊脚楼沿山势层层铺开。乘电瓶车分段游览，换上苗族服饰体验一回“苗家人”；'
        '还会用无人机为全团拍摄一段小视频，把苗寨的壮观与大家的笑容一起收进镜头。',
    '格鲁格桑十八寨（景中店）':
        '凯里一带的民族村寨景区，可感受黔东南少数民族的村寨风情，'
        '为接下来的荔波山水之旅添上一抹民族色彩。',
    '小七孔景区（电瓶车）':
        '荔波的“绿宝石”，属世界自然遗产“中国南方喀斯特”的一部分。乘电瓶车游览，'
        '一路是翡翠般透亮的湖水、层层跌落的叠水瀑布和长在水中的森林；景区因一座清代七孔古石桥而得名，美得像一幅水彩画。',
    '大七孔景区':
        '与小七孔相邻、风格却截然不同：这里更原始、更幽深。峡谷两侧峭壁高耸，河水碧绿清澈，'
        '标志性的“天生桥”横跨山间，气势磅礴，是喜欢自然野趣的游客心中的秘境。',
    '青云市集':
        '贵阳的特色街区市集，汇集各类黔味小吃摊档与本地小店，'
        '漫步其间，可以感受贵阳这座城市的日常生活气息。',
    '甲秀楼':
        '贵阳的城市名片，始建于明代万历年间，取“科甲挺秀”之意。三层飞檐的楼阁立于南明河中的巨石之上，'
        '由古朴的浮玉桥与岸边相连，楼影倒映在南明河上，古雅端庄。',
}

DAYS = [
    {
        'header': 'Day 1  新加坡→重庆→安顺    （午：夜郎风味｜晚：中式合菜）',
        'summary': (
            '多彩贵州之旅启程！由新加坡搭乘PN6426（0040/0540）抵达重庆，机场接机后转乘高铁前往安顺'
            '（暂定G9583  0911/1215）。随后游览龙宫景区（含电梯+游船），乘船探秘神奇的地下暗河溶洞。'
        ),
        'sights': ['龙宫景区（电梯+游船）'],
        'hotel': HOTEL_AS,
    },
    {
        'header': 'Day 2  安顺    （午：花开富贵｜晚：迎宾风味）',
        'summary': (
            '今天是与大瀑布的约会。早餐后前往黄果树瀑布景区（含环保车+往返扶梯），近距离感受飞瀑的磅礴气势；'
            '随后来到陡坡塘瀑布景区，寻找经典电视剧《西游记》里的熟悉画面。'
        ),
        'sights': ['黄果树瀑布景区（环保车+往返扶梯）', '陡坡塘瀑布景区'],
        'hotel': HOTEL_AS,
    },
    {
        'header': 'Day 3  安顺→黔西→贵阳    （午：中式合菜｜晚：黔菜风味）',
        'summary': (
            '今天前往以杜鹃花闻名的百里杜鹃。早餐后乘车前往黔西（约2H），游览百里杜鹃景区（含电瓶车）；'
            '之后乘车前往省会贵阳（约2H）。'
        ),
        'sights': ['百里杜鹃景区（电瓶车）'],
        'hotel': HOTEL_GY,
    },
    {
        'header': 'Day 4  贵阳→凯里    （午：苗家风味｜晚：酸汤鱼）',
        'summary': (
            '今天深入苗岭腹地。早餐后乘车前往凯里（约3H），走进西江千户苗寨（含四段电瓶车），'
            '换上苗服体验苗家风情，并由无人机为全团拍摄团体小视频，留下独一无二的旅行纪念。'
        ),
        'sights': ['西江千户苗寨（四段电瓶车+苗服体验+无人机拍摄团体小视频）'],
        'hotel': HOTEL_KL,
    },
    {
        'header': 'Day 5  凯里→荔波    （午：瑶家风味｜晚：石锅鱼）',
        'summary': (
            '从民族风情走进绿水青山。早餐后前往格鲁格桑十八寨（景中店）；随后乘车前往荔波（约3H），'
            '游览小七孔景区（含电瓶车），欣赏如翡翠般清澈的湖水与叠瀑。'
        ),
        'sights': ['格鲁格桑十八寨（景中店）', '小七孔景区（电瓶车）'],
        'hotel': HOTEL_LB,
    },
    {
        'header': 'Day 6  荔波→贵阳    （午：农家风味｜晚：黔菜风味）',
        'summary': (
            '今天在山水与城市之间切换。早餐后游览大七孔景区，探访原始幽深的峡谷秘境；随后乘车返回贵阳（约3.5H），'
            '漫步青云市集，再前往城市地标甲秀楼。'
        ),
        'sights': ['大七孔景区', '青云市集', '甲秀楼'],
        'hotel': HOTEL_GY,
    },
    {
        'header': 'Day 7  贵阳→重庆→新加坡    （午：重庆火锅）',
        'summary': (
            '早餐后乘高铁前往重庆（暂定D9932  0942/1152），午餐品尝地道的重庆火锅。'
            '餐后送往重庆机场，搭乘PN6425（1840/2340）返回新加坡，带着满满的回忆，回到温馨的家。'
        ),
        'sights': [],
        'hotel': None,
    },
]


def check():
    for day in DAYS:
        for name in day['sights']:
            if name not in SIGHT_INTROS:
                raise SystemExit(f'missing intro for {name}')
            n = len(re.findall(r'[\u4e00-\u9fff]', SIGHT_INTROS[name]))
            print(f'{n:3d}  {name}')


def reset_core_properties(doc):
    """模板壳会带来旧作者、旧日期；客户稿只留本行程标题。"""
    cp = doc.core_properties
    now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
    cp.title = TITLE
    cp.author = ''
    cp.last_modified_by = ''
    cp.comments = ''
    cp.keywords = ''
    cp.subject = ''
    cp.category = ''
    cp.revision = 1
    cp.created = now
    cp.modified = now


def strip_app_statistics(path):
    """docProps/app.xml 里的页数/字数/编辑时长是模板旧值，删掉让 Word 重新统计；core.xml 的 lastPrinted 也删掉。"""
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    tags = ['TotalTime', 'Pages', 'Words', 'Characters', 'Lines', 'Paragraphs',
            'CharactersWithSpaces', 'Company', 'Manager']
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == 'docProps/app.xml':
                for t in tags:
                    data = re.sub(f'<{t}>[^<]*</{t}>|<{t}/>'.encode(), b'', data)
            elif info.filename == 'docProps/core.xml':
                data = re.sub(rb'<cp:lastPrinted>[^<]*</cp:lastPrinted>', b'', data)
            z.writestr(info, data)


def build():
    check()
    g.SIGHT_INTROS = SIGHT_INTROS

    doc = Document(TEMPLATE)
    g.clear_body_keep_sectpr(doc)
    g.add_para(doc, TITLE, size=20, bold=True, color=g.C_TITLE,
               align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=3, line=1.2)
    g.add_para(doc, SUBTITLE, size=11, bold=True, color=g.C_SUB,
               align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, line=1.15)
    g.add_para(doc, FEATURES, size=11, bold=True, color=g.C_TEAL,
               before=2, after=6, line=1.15)
    for day in DAYS:
        start = len(doc.paragraphs)
        g.add_day_block(doc, day)
        # Day 条与概述不与后文分页
        for p in doc.paragraphs[start:start + 2]:
            p.paragraph_format.keep_with_next = True
    reset_core_properties(doc)
    doc.save(OUTPUT)
    strip_app_statistics(OUTPUT)
    print('Saved:', OUTPUT, os.path.getsize(OUTPUT))


if __name__ == '__main__':
    build()
