#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""English edition of 漫游成都6日（卧龙）— same structure/facts as gen_chengdu6.py, no pricing."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import generate_chongqing_sanxia8_voyage_style as g  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402

TEMPLATE = r'C:\Users\DELL\Downloads\12D11N_东北草原深度游_2026年6月.docx'
OUTPUT = r'C:\Users\DELL\Downloads\20261211漫游成都6日（卧龙）_行程_EN.docx'

FONT_LATIN = 'Arial'
FONT_EA = 'SimSun'


def set_run_font_en(run, name=g.FONT_CN, size_pt=11, bold=None, color=None):
    """Latin text in Arial, bracketed Chinese falls back to SimSun. Day bars keep Arial Unicode MS."""
    latin = name if name == g.FONT_DAY else FONT_LATIN
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), latin)
    rFonts.set(qn('w:hAnsi'), latin)
    rFonts.set(qn('w:cs'), latin)
    rFonts.set(qn('w:eastAsia'), FONT_EA if latin == FONT_LATIN else latin)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is True:
        run.font.bold = True
        if rPr.find(qn('w:bCs')) is None:
            rPr.append(OxmlElement('w:bCs'))
    if color is not None:
        run.font.color.rgb = color


g.set_run_font = set_run_font_en

TITLE = '6 Days 5 Nights · Leisurely Chengdu Discovery'
SUBTITLE = 'Singapore · Chengdu · Wolong · Dujiangyan'
FEATURES = 'Tour Features: Pure sightseeing · No shopping stops · No optional tours'

HOTEL_A = ('Chengdu Fengqi Museum Hotel (成都凤栖博物馆酒店) or Libai Hotel Chengdu Kuanzhai Alley & '
           'Wuhou Shrine (成都宽窄巷子武侯祠丽柏酒店) (4-Diamond)')
HOTEL_B = 'Langlizi Hotel, Dujiangyan Scenic Area (朗丽兹酒店·都江堰景区店) (4-Diamond)'
HOTEL_C = 'Jinglan Qingtang Hotel, Chunxi Road Taikoo Li (景澜青棠酒店·成都春熙路太古里店) (4-Diamond)'

S_WENSHU = 'Wenshu Monastery (文殊院)'
S_WENSHUFANG = 'Wenshufang Quarter (文殊坊)'
S_QINGYANG = 'Qingyang Palace (青羊宫)'
S_JINLI = 'Jinli Ancient Street (锦里)'
S_CHUNXI = 'Chunxi Road (春熙路)'
S_TAIKOO = 'Taikoo Li (太古里)'
S_WOLONG = 'Wolong Shenshuping Panda Base (卧龙神树坪基地) – incl. shuttle bus & audio guide'
S_DJY = 'Dujiangyan Irrigation System (都江堰) – incl. shuttle bus & headset'
S_NANQIAO = 'Nanqiao Bridge (南桥)'
S_BLUE = 'Blue Tears (蓝眼泪)'
S_PARK = "People's Park (人民公园)"
S_TEA = 'Heming Teahouse (鹤茗茶社) – tea included'
S_KUIXING = 'Kuixing Tower (魁星楼)'
S_EAST = 'Eastern Suburb Memory (东郊记忆)'
S_JIUYAN = 'Jiuyan Bridge (九眼桥)'

SIGHT_INTROS = {
    S_WENSHU:
        'A thousand-year-old temple hidden in the heart of the city, first built in the Sui Dynasty. '
        'Incense curls between red walls and golden-tiled roofs, old trees shade the halls, and the moment '
        "you step through the gate the city noise seems to fall away — a perfect first taste of Chengdu's "
        'calm, unhurried spirit.',
    S_WENSHUFANG:
        'A traditional Western Sichuan quarter right beside the monastery: grey-tiled roofs, white walls and '
        'carved wooden windows. The air carries the aroma of local snacks such as Zhong dumplings and sweet '
        'sesame rice balls, and around almost every corner sits an old teahouse — everyday Chengdu at its '
        'most authentic.',
    S_QINGYANG:
        "One of Southwest China's most celebrated Taoist temples, linked by legend to the sage Laozi and his "
        '"meeting at the Green Goat". Look for the finely carved dragon pillars of the Bagua Pavilion and the '
        'bronze goats rubbed to a shine by countless hands — touching one is said to bring good luck.',
    S_JINLI:
        'Right next to Wuhou Shrine, this old-style street is nicknamed "the First Street of Western Shu". '
        'Red lanterns, flagstone paths and Three Kingdoms-themed shop signs line the way, with sugar '
        'paintings, "three cannons" rice cakes and Sichuan opera masks at every turn; when night falls, '
        'the whole street glows.',
    S_CHUNXI:
        "Chengdu's busiest century-old shopping street, where trendy department stores stand beside "
        'time-honoured brands and the crowds never stop. Look up to spot the famous giant panda "climbing" '
        "the wall of Chengdu IFS — the best place to feel the city's energy.",
    S_TAIKOO:
        'An open-air quarter beside the ancient Daci Temple. Beneath sloping grey-tiled Sichuan-style roofs '
        "you'll find international labels and popular cafés; old and modern blend seamlessly here, and the "
        'lanes are especially charming once the lights come on.',
    S_WOLONG:
        'Tucked deep in the mountains of the Wolong National Nature Reserve, this panda base is wrapped in '
        "mist and bamboo-covered slopes — far closer to the pandas' wild home than any city zoo. Ride the "
        'shuttle bus in, listen through your audio guide, and watch the pandas munch bamboo, climb trees and '
        'tumble about.',
    S_DJY:
        'An engineering marvel built more than 2,000 years ago under Li Bing and his son. Without a single '
        'dam, it still controls floods, flushes away sand and irrigates the Chengdu Plain — the reason the '
        'region is called the "Land of Abundance". Now a UNESCO World Heritage Site; ride the shuttle bus, '
        'listen through your headset and stand beside the roaring Min River to marvel at ancient ingenuity.',
    S_NANQIAO:
        'The "front door" of old Dujiangyan town: a classical covered bridge with upturned eaves and painted '
        'beams, spanning the rushing river. As the evening lights come on, the bridge is mirrored in the '
        'water like a living painting.',
    S_BLUE:
        "Dujiangyan's most dreamlike night scene: lit from the banks, the rushing waters of the Min River's "
        'inner channel glow an ethereal blue, like a flowing ribbon of light — romantically nicknamed the '
        '"Blue Tears".',
    S_PARK:
        'To see real everyday Chengdu, come here: locals sip tea, play cards, sing and dance, and the famous '
        '"matchmaking corner" bustles with parents seeking partners for their children. The park\'s monument '
        'to the 1911 Railway Protection Movement quietly recalls a turning point in modern Chinese history.',
    S_TEA:
        "A long-established open-air teahouse by the lake in People's Park, filled with bamboo chairs. Cradle "
        'a bowl of lidded-cup tea, watch the locals "bai long men zhen" (chat for hours), and experience the '
        '"bashi" (comfortable and carefree) art of spending a whole afternoon over a single cup.',
    S_KUIXING:
        'A buzzing old-town neighbourhood marked by the tall Kuixing Tower. Its lanes hide Sichuan eateries '
        'serving chuan chuan skewers and mao cai alongside artsy cafés; under the neon at night, it is a '
        'favourite hangout for young locals.',
    S_EAST:
        'A creative park converted from a former state-owned electronic tube factory. Red-brick workshops, '
        'giant pipes and graffiti walls are everywhere; today it hosts live music, art exhibitions and pop-up '
        'markets — retro and brimming with creative energy.',
    S_JIUYAN:
        'A landmark stone arch bridge over the Jinjiang River, beautiful at night when its lights shimmer on '
        "the water. Bars line both banks and music drifts through the air — Chengdu's best-known nightlife "
        'district, revealing the lively, romantic side of this "slow city".',
}

DAYS = [
    {
        'header': '11Dec Day 1  Singapore → Chengdu    (B: X ｜ L: X ｜ D: X)',
        'summary': (
            'Depart Singapore for Chengdu, capital of the "Land of Abundance". Upon arrival at Chengdu '
            'Tianfu International Airport, you will be met and transferred to your hotel for check-in — '
            'your leisurely journey through Chengdu begins.'
        ),
        'sights': [],
        'hotel': HOTEL_A,
    },
    {
        'header': '12Dec Day 2  Chengdu    (B: Hotel ｜ L: Included ｜ D: Included)',
        'summary': (
            "Today we step into Chengdu's old-town soul. Begin with the calm of Wenshu Monastery and stroll "
            'the Sichuan-style lanes of Wenshufang; continue to Qingyang Palace to discover Taoist culture; '
            'then wander the Three Kingdoms-themed streets of Jinli; and finish on Chunxi Road and Taikoo Li, '
            'where Chengdu shows off its trendiest side.'
        ),
        'sights': [S_WENSHU, S_WENSHUFANG, S_QINGYANG, S_JINLI, S_CHUNXI, S_TAIKOO],
        'hotel': HOTEL_A,
    },
    {
        'header': '13Dec Day 3  Chengdu → Wolong → Dujiangyan    (B: Hotel ｜ L: Included ｜ D: Included)',
        'summary': (
            'Panda day is finally here! Head to Wolong Shenshuping Panda Base (incl. shuttle bus & audio '
            'guide) to see giant pandas amid mountain bamboo forests. Next, visit the Dujiangyan Irrigation '
            'System (incl. shuttle bus & headset), an ancient waterworks still at work after more than 2,000 '
            'years. Then stroll over Nanqiao Bridge and, after dark, take in the dreamy "Blue Tears".'
        ),
        'sights': [S_WOLONG, S_DJY, S_NANQIAO, S_BLUE],
        'hotel': HOTEL_B,
    },
    {
        'header': '14Dec Day 4  Dujiangyan → Chengdu    (B: Hotel ｜ L: Included ｜ D: Included)',
        'summary': (
            "Back to Chengdu for a taste of true local slow living: watch everyday life in People's Park and "
            'sip a bowl of lidded-cup tea at Heming Teahouse (tea included); soak up the street-level buzz '
            'around Kuixing Tower; hunt for art and creativity among the old factory halls of Eastern Suburb '
            "Memory; and end at Jiuyan Bridge, enjoying Chengdu's evening glow along the Jinjiang River."
        ),
        'sights': [S_PARK, S_TEA, S_KUIXING, S_EAST, S_JIUYAN],
        'hotel': HOTEL_C,
    },
    {
        'header': '15Dec Day 5  Chengdu    (B: Hotel ｜ L: X ｜ D: X)',
        'summary': (
            'Free day at leisure. No schedule to keep — sleep in and explore Chengdu at your own pace. Your '
            "hotel sits in the Chunxi Road / Taikoo Li area, so the city's liveliest streets are right at "
            'your doorstep.'
        ),
        'sights': [],
        'hotel': HOTEL_C,
    },
    {
        'header': '16Dec Day 6  Chengdu → Singapore    (B: Hotel ｜ L: X ｜ D: X)',
        'summary': (
            'After breakfast, transfer to Chengdu Tianfu International Airport for your flight home to '
            'Singapore, carrying wonderful memories of Chengdu.'
        ),
        'sights': [],
        'hotel': None,
    },
]


def check_intros():
    for day in DAYS:
        for name in day['sights']:
            if name not in SIGHT_INTROS:
                raise SystemExit(f'missing intro for {name}')


def add_day_block_en(doc, day):
    g.add_para(doc, day['header'], size=11, bold=True, color=g.C_WHITE, font=g.FONT_DAY,
               before=10, after=4, left=2, shade=g.C_DAY_BG, line=1.1)
    g.add_para(doc, day['summary'], size=10.5, color=g.C_BODY, justify=True,
               before=3, after=3, line=1.25)
    for name in day['sights']:
        g.add_sight_para(doc, name, ' ' + SIGHT_INTROS[name])
    if day['hotel']:
        g.add_para(doc, f'Stay: {day["hotel"]}', size=10.5, bold=True, color=g.C_TITLE,
                   before=4, after=2, left=2, line=1.1)


def build():
    check_intros()
    doc = Document(TEMPLATE)
    g.clear_body_keep_sectpr(doc)
    g.add_para(doc, TITLE, size=20, bold=True, color=g.C_TITLE,
               align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=3, line=1.2)
    g.add_para(doc, SUBTITLE, size=11, bold=True, color=g.C_SUB,
               align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, line=1.15)
    g.add_para(doc, FEATURES, size=11, bold=True, color=g.C_TEAL,
               before=2, after=6, line=1.15)
    for day in DAYS:
        start = len(doc.paragraphs)
        add_day_block_en(doc, day)
        for p in doc.paragraphs[start:start + 2]:
            p.paragraph_format.keep_with_next = True
    doc.save(OUTPUT)
    print('Saved:', OUTPUT, os.path.getsize(OUTPUT))


if __name__ == '__main__':
    build()
