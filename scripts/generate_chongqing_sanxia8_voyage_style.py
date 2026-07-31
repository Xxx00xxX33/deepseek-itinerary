#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
学习并复刻：
  C:\\Users\\DELL\\Downloads\\重庆+长江三峡揽月8D7N上水_行程 (2).docx

- 前部 / Day 条 / 概述 / 宿： 完全按 (2) 的结构与样式
- 仅景点【】后的描述改为「地点介绍」40–50 汉字（不计标点）
- 不含报价段
"""
from __future__ import annotations

import os
import re
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

# 修改好的版本：样式与正文骨架来源
SOURCE = r'C:\Users\DELL\Downloads\重庆+长江三峡揽月8D7N上水_行程 (2).docx'

OUTPUTS = [
    r'C:\Users\DELL\Downloads\重庆+长江三峡揽月8D7N上水_行程.docx',
    r'C:\Users\DELL\Downloads\重庆+长江三峡揽月8D7N上水_行程_无报价.docx',
    r'C:\Users\DELL\Downloads\Chongqing_Yangtze_Lanyue_8D7N_Upstream.docx',
]

# 与 (2) 一致
C_TITLE = RGBColor(0x1A, 0x6B, 0x5E)
C_SUB = RGBColor(0x55, 0x55, 0x55)
C_TEAL = RGBColor(0x00, 0x80, 0x80)
C_BODY = RGBColor(0x00, 0x00, 0x00)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_DAY_BG = '2E8B7A'  # (2) Day 条底色

FONT_CN = 'SimSun'           # (2) 正文
FONT_DAY = 'Arial Unicode MS'  # (2) Day 条

CJK_RE = re.compile(r'[\u4e00-\u9fff]')


def cjk_count(text: str) -> int:
    return len(CJK_RE.findall(text))


def set_run_font(run, name=FONT_CN, size_pt=11, bold=None, color=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:cs'), name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is True:
        run.font.bold = True
        if rPr.find(qn('w:bCs')) is None:
            rPr.append(OxmlElement('w:bCs'))
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_shading(paragraph, fill_hex: str):
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn('w:shd')):
        pPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    pPr.append(shd)


def set_spacing(paragraph, before=None, after=None, left=None, line=None):
    pf = paragraph.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if left is not None:
        pf.left_indent = Pt(left)
    if line is not None:
        pf.line_spacing = line


def clear_body_keep_sectpr(doc: Document):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)


def add_para(doc, text, *, size=11, bold=False, color=C_BODY, font=FONT_CN,
             align=None, before=0, after=4, left=None, justify=False,
             shade=None, line=1.15):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align is not None:
        p.alignment = align
    set_spacing(p, before=before, after=after, left=left, line=line)
    if shade:
        set_paragraph_shading(p, shade)
    if text:
        run = p.add_run(text)
        set_run_font(run, name=font, size_pt=size, bold=bold if bold else None, color=color)
    return p


def add_sight_para(doc, name: str, desc: str):
    """与 (2) 一致：【名】1A6B5E 加粗 + 正文黑。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_spacing(p, before=2, after=3, line=1.15)
    r1 = p.add_run(f'【{name}】')
    set_run_font(r1, name=FONT_CN, size_pt=10.5, bold=True, color=C_TITLE)
    r2 = p.add_run(desc)
    set_run_font(r2, name=FONT_CN, size_pt=10.5, color=C_BODY)
    return p


def add_run_text(p, text, *, size=11, bold=False, color=C_BODY, font=FONT_CN):
    run = p.add_run(text)
    set_run_font(run, name=font, size_pt=size, bold=bold if bold else None, color=color)
    return run


# ---------- 前部（与 (2) 一致） ----------
TITLE = '8D7N重庆+壮丽三峡长江行'
SUBTITLE = '新加坡·重庆·宜昌·长江三峡（长江行揽月号）'
SPECIAL = '特别安排：2026年4月下水、一价全包的长江智慧游轮─《长江行揽月号》'

HIGHLIGHTS = [
    '【一价全包】全程无任何上岸自费景点及隐形消费，打造真正“一价全包”的游轮体验',
    '【两坝三峡】体验三峡大坝升船机、葛洲坝船闸、白天过全程三峡，看大国重器，享峡江之美',
    '【智能创新】长江首艘搭载Harmony OS 鸿蒙智能舱房系统的智慧游轮，提供前沿科技的便捷体验',
    '【WiFi畅联】全船免费，通过技术手段将信号覆盖半径从近岸5公里拓展至50公里，优化卡顿痛点',
    '【港韵雅宴】由“食神”戴龙亲传弟子、中国烹饪大师张建军领衔，带来一日一主题的味觉盛宴',
    '【沉浸演艺】长江新体验，专职NPC陪玩，体验真正的度假之旅',
]

SHIP_INFO = '全船载客：650人'
FEATURES = '行程特色：纯玩无购物团；无额外必加景点费用； 全程品质酒店 + 全新长江智慧游轮'
FOOD = '各地美食：重庆烤肉自助、重庆特色火锅、船长欢送晚宴'
GIFTS = '揽月号赠送项目：换乘豪华观光船【船过葛洲坝船闸】'

# 景点地点介绍（唯一不照抄 (2) 的部分）
SIGHT_INTROS = {
    '苏家坝立交观景平台、弹子石老街':
        '苏家坝立交是重庆立体交通的代表，观景平台可俯瞰多层道路与江景。弹子石老街保留开埠时期街巷与老建筑。',
    '李子坝轻轨穿楼（含乘坐1-2站）':
        '李子坝轻轨站因轨道从居民楼中间穿过而闻名。行程含乘坐一至两站，可近距离感受山城独特交通景观。',
    '解放碑步行街、八一路好吃街、洪崖洞夜景':
        '解放碑是重庆地标纪念碑，周边为繁华步行街。八一路集中本地小吃。洪崖洞以多层吊脚楼建筑和夜景灯光著称。',
    '重庆→宜昌高铁':
        '高铁连接重庆与宜昌，穿行西南山地。参考车次约九时四十一分开车、十四时二十一分抵达，车程约四个半小时。',
    '宜昌登船':
        '宜昌是长江三峡东端门户，登船点在三峡游客中心。旅客在十五时至十五时三十分之间完成手续，登上游轮。',
    '船过葛洲坝船闸·游西陵峡东段（揽月*赠送项目）':
        '葛洲坝是长江上大型水利枢纽。观光船经船闸升降并游览西陵峡东段，可直观感受水位落差与峡谷江景。',
    '三峡大坝、升船机':
        '三峡大坝是世界规模最大的水利枢纽。升船机以垂直升降方式运送船舶过坝；停运时改参观博物馆与截流纪念园。',
    '西陵峡西段':
        '西陵峡是长江三峡中最长的一段。西段两岸山势连绵、江面开阔，游轮白日航行可连续观赏峡江自然风光。',
    '船长欢迎酒会':
        '船长欢迎酒会是游轮开航后的船上活动。由船长致辞并介绍航行与服务安排，旅客可借此熟悉游轮环境。',
    '巴东雨仙谷（原无源洞）/巫山神女溪':
        '雨仙谷原无源洞以溶洞与光影景观见长，旺季夜游；神女溪是巫山支流峡谷，淡季改上午游览，感受峡溪风光。',
    '巫峡&瞿塘峡':
        '巫峡以幽深秀丽闻名，瞿塘峡以江窄山高、气势雄峻著称。两段同为长江三峡核心峡谷，白日乘船经过。',
    '白帝城 / 三峡之巅（二选一）':
        '白帝城坐落长江北岸，与三国历史及唐宋诗篇相关。三峡之巅位于瞿塘峡口，可俯瞰峡谷群峰与长江水道。',
    '游轮文艺表演':
        '游轮文艺表演在船上演出厅举行。内容由游轮统一编排，多为歌舞与主题演出，供旅客在航程中晚间观赏。',
    '在船文化活动':
        '在船文化活动由游轮在航行日上午安排。旅客可参与讲座、体验或休闲项目，并利用时间熟悉船上各项设施。',
    '白鹤梁+涪陵榨菜馆 / 816工程遗址（二选一）':
        '白鹤梁以水下水文题刻闻名，并配榨菜历史展示。八一六为大型地下人工洞体工程；不可抗力时改丰都鬼城。',
    '游轮欢送晚会':
        '欢送晚会安排在游轮航程接近结束的夜晚。作为船上告别活动，旅客可与船员及其他旅客共同参与观赏。',
    '朝天门码头下船':
        '朝天门位于长江与嘉陵江交汇处，是重庆传统水上门户。游轮在此靠岸下船后，转入市区陆地行程继续游览。',
    '磁器口古镇、观音桥步行街、北仓文创园、光环购物中心':
        '磁器口是依嘉陵江的传统古镇，保留石板街与老店铺。观音桥、北仓与光环则是重庆现代商业文创与购物街区。',
}

# Day 数据完全按 (2)：含日期前缀、餐食写法、概述、酒店
DAYS = [
    {
        'header': '06Dec Day 1  新加坡→重庆    （午：X｜晚：X）',
        'summary': '齐聚樟宜机场，飞往8D魔都 – 重庆。专人接机，入住酒店。',
        'sights': [],
        'hotel': '南坪嘉发希尔顿逸林酒店或同级（五钻）',
    },
    {
        'header': '07Dec Day 2  重庆    （午：重庆风味｜晚：中式合菜）',
        'summary': (
            '游览苏家坝立交观景平台、弹子石老街、李子坝轻轨穿楼体验（含乘坐1-2站）、'
            '解放碑步行街、八一路好吃街、洪崖洞夜景。'
        ),
        'sights': [
            '苏家坝立交观景平台、弹子石老街',
            '李子坝轻轨穿楼（含乘坐1-2站）',
            '解放碑步行街、八一路好吃街、洪崖洞夜景',
        ],
        'hotel': '南坪嘉发希尔顿逸林酒店或同级（五钻）',
    },
    {
        'header': '08Dec Day 3  重庆→宜昌（周二登船）    （午：高铁午餐｜晚：登船晚餐）',
        'summary': (
            '乘坐高铁前往宜昌（参考车次0941-1421）；'
            '接站，送宜昌三峡游客中心登船（于15.00-15.30之间抵达）；'
            '参加揽月*赠送项目：换乘豪华观光船【船过葛洲坝船闸·游西陵峡东段】'
            '—抵达三斗坪车送秭归港登船。'
        ),
        'sights': [
            '重庆→宜昌高铁',
            '宜昌登船',
            '船过葛洲坝船闸·游西陵峡东段（揽月*赠送项目）',
        ],
        'hotel': '长江行揽月号',
    },
    {
        'header': '09Dec Day 4  三峡    （午：游船自助｜晚：游船自助）',
        'summary': (
            '上午参观游览【三峡大坝】，抵达三斗坪换乘游船过【升船机】'
            '（重要提示：因天气、水文、航道等原因及三峡升船机禁航停运期间，'
            '调整为参观三峡工程博物馆&截流纪念园后车送秭归港登船）；'
            '午餐后游轮航行经过【西陵峡西段】—欣赏长江山水的雄浑与灵秀；'
            '晚餐前船长欢迎酒会；'
            '晚上夜游【巴东雨仙谷（原无源洞）】—身临光影画卷，寻境别有洞天'
            '（重要提示：4月-10月参观巴东雨仙谷（原无源洞），'
            '11月-次年3月安排次日上午游览巫山神女溪）。'
        ),
        'sights': [
            '三峡大坝、升船机',
            '西陵峡西段',
            '船长欢迎酒会',
            '巴东雨仙谷（原无源洞）/巫山神女溪',
        ],
        'hotel': '长江行揽月号',
    },
    {
        'header': 'Day 5  三峡    （午：游船自助｜晚：游船自助）',
        'summary': (
            '上午游轮航行经过【巫峡&瞿塘峡】；'
            '下午上岸游览（二选一）：'
            'A线：【白帝城】—体验三国古迹与诗城的千古遗风；'
            'B线：【三峡之巅】—俯瞰峡谷群峰与长江奔流；'
            '晚上游轮文艺表演。午/晚：游船自助。'
        ),
        'sights': [
            '巫峡&瞿塘峡',
            '白帝城 / 三峡之巅（二选一）',
            '游轮文艺表演',
        ],
        'hotel': '长江行揽月号',
    },
    {
        'header': 'Day 6  三峡→重庆    （午：游船自助｜晚：游船自助）',
        'summary': (
            '上午在船文化活动；'
            '下午上岸游览（二选一）：'
            'A线：【白鹤梁水下博物馆+涪陵榨菜历史记忆馆】—水下千年对话，舌尖百年传承；'
            'B线：【816工程遗址】—探秘全球最大人工洞体'
            '（重要提示：若因天气、水文、航道等不可抗力情况停靠丰都港，上岸游览调整为丰都鬼城）；'
            '晚上游轮欢送晚会。午/晚：游船自助。'
        ),
        'sights': [
            '在船文化活动',
            '白鹤梁+涪陵榨菜馆 / 816工程遗址（二选一）',
            '游轮欢送晚会',
        ],
        'hotel': '长江行揽月（超五星）',
    },
    {
        'header': 'Day 7  重庆    （午：融合菜｜晚：重庆特色火锅风味）',
        'summary': '朝天门码头下船，游览磁器口古镇、观音桥步行街、北仓文创园、光环购物中心。',
        'sights': [
            '朝天门码头下船',
            '磁器口古镇、观音桥步行街、北仓文创园、光环购物中心',
        ],
        'hotel': '南坪嘉发希尔顿逸林酒店或同级（五钻）',
    },
    {
        'header': 'Day 8  重庆→新加坡    （午：×｜晚：×）',
        'summary': '早餐后，专人送往机场，返回家园 – 新加坡',
        'sights': [],
        'hotel': None,
    },
]


def validate_sight_intros():
    bad = []
    for name, desc in SIGHT_INTROS.items():
        n = cjk_count(desc)
        if not (40 <= n <= 50):
            bad.append((name, n, desc))
    if bad:
        for item in bad:
            print('INTRO_LEN_FAIL:', item)
        raise SystemExit(f'{len(bad)} sight intro(s) not in 40-50 CJK chars')
    # every day sight must have intro
    for day in DAYS:
        for name in day['sights']:
            if name not in SIGHT_INTROS:
                raise SystemExit(f'missing intro for {name}')
    print('Sight intro CJK counts OK (40-50, punctuation excluded)')


def add_front_matter(doc: Document):
    """前部样式对齐 (2)：标题 20pt 1A6B5E；副标 11pt 555555；亮点 11pt 008080。"""
    add_para(doc, TITLE, size=20, bold=True, color=C_TITLE,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=3, line=1.2)

    add_para(doc, SUBTITLE, size=11, bold=True, color=C_SUB,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, line=1.15)

    # 特别安排：1A6B5E 加粗 ~10.5pt，两端对齐
    add_para(doc, SPECIAL, size=10.5, bold=True, color=C_TITLE,
             justify=True, before=2, after=4, line=1.15)

    for line in HIGHLIGHTS:
        add_para(doc, line, size=11, color=C_TEAL, justify=True,
                 before=0, after=2, line=1.15)

    add_para(doc, SHIP_INFO, size=11, color=C_TEAL, justify=True,
             before=2, after=2, line=1.15)

    # 行程特色 / 美食 / 赠送：008080 加粗（与 (2) 一致）
    add_para(doc, FEATURES, size=11, bold=True, color=C_TEAL,
             before=2, after=2, line=1.15)
    add_para(doc, FOOD, size=11, bold=True, color=C_TEAL,
             before=1, after=2, line=1.15)
    add_para(doc, GIFTS, size=11, bold=True, color=C_TEAL,
             before=1, after=6, line=1.15)


def add_day_block(doc: Document, day: dict):
    # Day 条：(2) = 白字 + 2E8B7A 底 + Arial Unicode MS + before 10 / after 4
    add_para(
        doc, day['header'],
        size=11, bold=True, color=C_WHITE, font=FONT_DAY,
        before=10, after=4, left=2, shade=C_DAY_BG, line=1.1,
    )

    # 概述：两端对齐，SimSun 10.5，黑
    add_para(
        doc, day['summary'],
        size=10.5, color=C_BODY, justify=True,
        before=3, after=3, line=1.25,
    )

    for name in day['sights']:
        add_sight_para(doc, name, SIGHT_INTROS[name])

    if day['hotel']:
        add_para(
            doc, f'宿：{day["hotel"]}',
            size=10.5, bold=True, color=C_TITLE,
            before=4, after=2, left=2, line=1.1,
        )


def build():
    validate_sight_intros()

    if not os.path.isfile(SOURCE):
        raise FileNotFoundError(SOURCE)

    # 用 (2) 作壳：保留页眉 logo / section / 主题
    doc = Document(SOURCE)
    clear_body_keep_sectpr(doc)

    add_front_matter(doc)
    for day in DAYS:
        add_day_block(doc, day)

    saved = []
    errors = []
    for path in OUTPUTS:
        try:
            if not saved:
                doc.save(path)
            else:
                shutil.copy2(saved[0], path)
            print('Saved:', path, os.path.getsize(path))
            saved.append(path)
        except PermissionError:
            print('LOCKED:', path)
            errors.append(path)

    if not saved:
        raise SystemExit('No output saved; all targets locked')
    if errors:
        print('Some targets locked:', errors)


if __name__ == '__main__':
    build()
